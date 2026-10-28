"""The workbook must say exactly what the Python side says.

Every KPI formula in the workbook is evaluated with an independent engine (the
`formulas` package) and compared cell by cell with targets.py.
"""
import re

import formulas
import numpy as np
import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from kpi_kit.facts import GRAINS
from kpi_kit.kpis import KPIS
from kpi_kit.workbook import FIRST_ROW, SECTION_START, _sheet_names


@pytest.fixture(scope="module")
def evaluated(outputs):
    path = outputs / "kpi-report.xlsx"
    solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
    values = {}
    for key, value in solution.items():
        match = re.match(r"'\[[^\]]+\](.+)'!([A-Z]+\d+)$", key)
        if match:
            values[(match.group(1).upper(), match.group(2))] = value.value[0, 0] if hasattr(value, "value") else value
    return path, values


def test_no_formula_errors(evaluated):
    _, values = evaluated
    errors = [(s, c) for (s, c), v in values.items() if isinstance(v, formulas.functions.XlError)]
    assert errors == []


def test_kpi_tabs_match_python(evaluated, periods, results):
    _, values = evaluated
    mismatches = []
    for grain in GRAINS:
        tab = _sheet_names(grain)[1].upper()
        for j, kpi in enumerate(KPIS):
            letters = [get_column_letter(2 + 4 * j + x) for x in range(4)]
            for i, key in enumerate(periods[grain].table.index):
                for field, col in zip(["actual", "target", "perf", "band"], letters):
                    want = results[grain][kpi.id].at[key, field]
                    got = values.get((tab, f"{col}{FIRST_ROW + i}"))
                    got = None if got in ("", None) else float(got)
                    if not np.isfinite(want):
                        if got is not None:
                            mismatches.append((grain, kpi.id, key, field, want, got))
                    elif got is None or abs(got - want) > 1e-9 * max(1.0, abs(want)):
                        mismatches.append((grain, kpi.id, key, field, want, got))
    assert mismatches == []


def test_report_tabs_show_python_statuses(evaluated, results):
    path, values = evaluated
    wb = openpyxl.load_workbook(path)
    labels = {k.label: k for k in KPIS}
    for grain in GRAINS:
        name = _sheet_names(grain)[0]
        key = results[grain]["net_rev"].dropna(subset=["target"]).index[-1]
        seen = 0
        for (cell,) in wb[name].iter_rows(min_row=SECTION_START, min_col=2, max_col=2):
            kpi = labels.get(cell.value)
            if kpi:
                assert values[(name.upper(), f"F{cell.row}")] == results[grain][kpi.id].at[key, "status"]
                seen += 1
        assert seen == len(KPIS)
