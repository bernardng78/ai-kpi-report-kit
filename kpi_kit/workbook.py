"""Write the KPI workbook: report tabs on top of formula-driven KPI and data tabs.

The brand's numbers per period are the only typed-in values (blue). Everything
derived from them is a formula, so anyone can click a target and see how it
was built: derived columns on the data tabs, every actual, target, miss and
band on the KPI tabs, and every figure on the report tabs. The actions are the
one exception: they are text written by the rules in actions.py, looked up by
period.
"""
from __future__ import annotations

import math
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import design as d
from .clean import DataChecks
from .config import (BAND_CAP, BAND_FLOOR, DEFAULT_BAND, GRAIN_LABEL, GROWTH_MAX, GROWTH_MIN, WINDOWS,
                     Config)
from .facts import GRAINS, PeriodData
from .kpis import BASE_COLUMNS, KPIS, SECTIONS
from .targets import first_target_index

FIRST_ROW = 3            # data and KPI tabs: two header rows, then one row per period
SETTING_ROW = {"growth_min": 4, "growth_max": 5, "band_floor": 6, "band_cap": 7, "default_band": 8}
CHART_PERIODS = {"week": 13, "month": 13, "quarter": 5}
CHART_FIRST_COL = {"week": 1, "month": 8, "quarter": 15}  # each report's block on the Chart data tab
SWING_FORMAT = '"±"0.0%'
PERIOD_WORD = {"week": "Week", "month": "Month", "quarter": "Quarter"}
ROW_CELL = "$X$5"        # report tabs: position of the chosen period in the KPI tab
SECTION_START = 24       # report tabs: first section header row
NOTES = {
    "aov": "Gross revenue ÷ orders",
    "refund_rate": "Lower is better",
    "cr": "Shopify orders ÷ GA4 sessions",
    "spend": "Compared with trajectory, not judged",
    "mer": "Net revenue ÷ ad spend",
    "ncac": "Lower is better",
    "roas_meta": "Platform-reported: budget splits only",
    "roas_google": "Platform-reported: budget splits only",
    "roas_tiktok": "Platform-reported: budget splits only",
    "repeat_90d": "Customers who first bought ~90 days earlier",
    "cm": "After COGS, fulfilment, fees and ad spend",
}


def _font(bold: bool = False, size: int = 10, color: str = d.INK, italic: bool = False) -> Font:
    return Font(name=d.FONT, bold=bold, size=size, color=color, italic=italic)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


HAIR = Side(style="thin", color=d.GRID)
BOTTOM = Border(bottom=HAIR)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def _q(name: str) -> str:
    return f"'{name}'"


def _sheet_names(grain: str) -> tuple[str, str, str]:
    label = GRAIN_LABEL[grain]
    return f"{label} report", f"{label} KPIs", f"{label} data"


def _style_header(cells, fill: str | None = d.PANEL) -> None:
    for cell in cells:
        cell.font = _font(bold=True)
        cell.border = BOTTOM
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if fill:
            cell.fill = _fill(fill)


def build_workbook(path: Path, cfg: Config, periods: dict[str, PeriodData], actions: pd.DataFrame,
                   checks: DataChecks, defaults: dict[str, str]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    order = ["Read me"] + [_sheet_names(g)[0] for g in GRAINS] + [_sheet_names(g)[1] for g in GRAINS] \
        + [_sheet_names(g)[2] for g in GRAINS] + ["Chart data", "Actions", "KPI definitions", "Data checks", "Settings"]
    for name in order:
        wb.create_sheet(name)

    _settings(wb["Settings"])
    action_rows = _actions(wb["Actions"], actions)
    for grain in GRAINS:
        data = periods[grain]
        report, kpi_tab, data_tab = _sheet_names(grain)
        cols = _data_sheet(wb[data_tab], grain, data)
        kpi_cols = _kpi_sheet(wb[kpi_tab], grain, data, cols)
        _report_sheet(wb[report], grain, data, kpi_cols, action_rows, defaults[grain], cfg)
    _definitions(wb["KPI definitions"])
    _checks(wb["Data checks"], checks)
    _read_me(wb["Read me"], cfg, defaults)

    wb.active = wb.sheetnames.index(_sheet_names("week")[0])
    wb.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# --- inputs ---------------------------------------------------------------------

def _settings(ws) -> None:
    ws["A1"] = "Settings"
    ws["A1"].font = _font(bold=True, size=14)
    ws["A2"] = "The KPI tabs read these cells. Change a yellow cell and every target and status updates."
    ws["A2"].font = _font(color=d.INK_2)
    for j, header in enumerate(["Setting", "Value", "What it does"], start=1):
        ws.cell(3, j, header)
    _style_header(ws[3][:3])
    rows = [
        ("growth_min", "Growth factor floor", GROWTH_MIN, "0.00", "Lowest growth a target can assume (-33%)."),
        ("growth_max", "Growth factor cap", GROWTH_MAX, "0.00", "Highest growth a target can assume (+50%)."),
        ("band_floor", "Band floor", BAND_FLOOR, "0%", "Narrowest band: misses under half of it are always on track."),
        ("band_cap", "Band cap", BAND_CAP, "0%", "Widest band, for very noisy KPIs."),
        ("default_band", "Default band", DEFAULT_BAND, "0%", "Used until a KPI has a full window of past misses."),
    ]
    for key, label, value, fmt, note in rows:
        r = SETTING_ROW[key]
        ws.cell(r, 1, label).font = _font()
        cell = ws.cell(r, 2, value)
        cell.font = _font(color=d.INPUT_BLUE)
        cell.fill = _fill(d.SETTING_FILL)
        cell.number_format = fmt
        ws.cell(r, 3, note).font = _font(color=d.INK_2)

    ws["A11"] = "Target windows (built into the KPI formulas; change them in config.py and rebuild)"
    ws["A11"].font = _font(bold=True)
    for j, header in enumerate(["Grain", "Periods per year", "Anchor", "Growth window", "Band window"], start=1):
        ws.cell(12, j, header)
    _style_header(ws[12][:5])
    anchor_text = {0: "same period last year", 1: "same week last year ±1 week"}
    for i, grain in enumerate(GRAINS, start=13):
        w = WINDOWS[grain]
        values = [GRAIN_LABEL[grain], w["lag"], anchor_text[w["anchor"]], f"last {w['trailing']} vs a year earlier",
                  f"previous {w['band']}"]
        for j, v in enumerate(values, start=1):
            ws.cell(i, j, v).font = _font()
    for letter, width in zip("ABCDE", [24, 16, 30, 60, 16]):
        ws.column_dimensions[letter].width = width


def _actions(ws, actions: pd.DataFrame) -> int:
    headers = ["Lookup key", "Grain", "Period", "Section", "No.", "Action"]
    for j, h in enumerate(headers, start=1):
        ws.cell(1, j, h)
    _style_header(ws[1][:6])
    for i, row in enumerate(actions.itertuples(index=False), start=2):
        values = [f"{row.grain}|{row.key}|{row.section}|{row.n}", row.grain, row.key, row.section, row.n, row.text]
        for j, v in enumerate(values, start=1):
            cell = ws.cell(i, j, v)
            cell.font = _font()
        ws.cell(i, 6).alignment = WRAP_TOP
    for letter, width in zip("ABCDEF", [34, 9, 11, 18, 4, 120]):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    return len(actions) + 1  # last row


def _data_sheet(ws, grain: str, data: PeriodData) -> dict[str, str]:
    """Write one row per period; return base column id -> column letter."""
    ws["A1"] = (f"{GRAIN_LABEL[grain]} data: the brand's own numbers. Blue cells come from the exports; "
                f"black cells are formulas.")
    ws["A1"].font = _font(bold=True)
    headers = [("key", "Period"), ("start", "First day"), ("end", "Last day")] + \
              [(c, header) for c, header, _, _ in BASE_COLUMNS]
    cols = {cid: get_column_letter(j) for j, (cid, _) in enumerate(headers, start=1)}
    for j, (_, header) in enumerate(headers, start=1):
        ws.cell(2, j, header)
    _style_header(ws[2][:len(headers)])
    for i, (key, row) in enumerate(data.table.iterrows()):
        r = FIRST_ROW + i
        ws.cell(r, 1, key).font = _font()
        for j, col in [(2, "start"), (3, "end")]:
            cell = ws.cell(r, j, row[col].to_pydatetime())
            cell.number_format = "d mmm yyyy"
            cell.font = _font()
        for cid, _, fmt, formula in BASE_COLUMNS:
            cell = ws[f"{cols[cid]}{r}"]
            if formula:
                cell.value = "=" + formula.format(**{k: f"{v}{r}" for k, v in cols.items()})
                cell.font = _font()
            else:
                value = row[cid]
                if pd.notna(value):
                    cell.value = int(round(value)) if fmt == "int" else round(float(value), 2)
                cell.font = _font(color=d.INPUT_BLUE)
            cell.number_format = d.NUMBER_FORMAT[fmt]
    ws.column_dimensions["A"].width = 10
    for j in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 13
    ws.row_dimensions[2].height = 42
    ws.freeze_panes = "D3"
    return cols


# --- calculations ---------------------------------------------------------------

def _kpi_sheet(ws, grain: str, data: PeriodData, data_cols: dict[str, str]) -> dict[str, dict[str, str]]:
    """Actual, target, miss and band for every KPI and period; return KPI id -> column letters."""
    data_tab = _q(_sheet_names(grain)[2])
    w = WINDOWS[grain]
    lag, half, trailing, band_n = w["lag"], w["anchor"], w["trailing"], w["band"]
    n = len(data.table)
    s = {key: f"Settings!$B${row}" for key, row in SETTING_ROW.items()}

    def rng(col: str, first: int, last: int) -> str:
        return f"{data_tab}!{col}{first}:{col}{last}"

    def clamp(expr: str) -> str:
        return f"MIN({s['growth_max']},MAX({s['growth_min']},{expr}))"

    ws["A2"] = "Period"
    _style_header([ws["A1"], ws["A2"]])
    out = {}
    for j, kpi in enumerate(KPIS):
        c0 = 2 + 4 * j
        act, tgt, perf, band = (get_column_letter(c0 + x) for x in range(4))
        out[kpi.id] = {"actual": act, "target": tgt, "perf": perf, "band": band}
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + 3)
        ws.cell(1, c0, kpi.label)
        for x, header in enumerate(["Actual", "Target", "vs target", "Band"]):
            ws.cell(2, c0 + x, header)
        _style_header([ws.cell(1, c0 + x) for x in range(4)] + [ws.cell(2, c0 + x) for x in range(4)])
        num = data_cols[kpi.num]
        den = data_cols[kpi.den] if kpi.den else None
        first = first_target_index(data.table, kpi, grain)
        for i in range(n):
            r = FIRST_ROW + i
            if j == 0:
                ws[f"A{r}"] = f"={data_tab}!A{r}"
                ws[f"A{r}"].font = _font()
            if den:
                ws[f"{act}{r}"] = (f'=IF(AND(ISNUMBER({data_tab}!{num}{r}),ISNUMBER({data_tab}!{den}{r})),'
                                   f'IFERROR({data_tab}!{num}{r}/{data_tab}!{den}{r},""),"")')
            else:
                ws[f"{act}{r}"] = f'=IF(ISNUMBER({data_tab}!{num}{r}),{data_tab}!{num}{r},"")'
            if i >= first:
                a1, a2 = r - lag - half, r - lag + half
                t1, t2 = r - trailing, r - 1
                l1, l2 = r - lag - trailing, r - lag - 1
                if den:
                    growth = (f"(SUM({rng(num, t1, t2)})/SUM({rng(den, t1, t2)}))"
                              f"/(SUM({rng(num, l1, l2)})/SUM({rng(den, l1, l2)}))")
                    ws[f"{tgt}{r}"] = (f'=IFERROR(SUM({rng(num, a1, a2)})/SUM({rng(den, a1, a2)})'
                                       f'*{clamp(growth)},"")')
                else:
                    growth = f"SUM({rng(num, t1, t2)})/SUM({rng(num, l1, l2)})"
                    ws[f"{tgt}{r}"] = f'=IFERROR(AVERAGE({rng(num, a1, a2)})*{clamp(growth)},"")'
                miss = f"1-{act}{r}/{tgt}{r}" if kpi.direction == "down" else f"{act}{r}/{tgt}{r}-1"
                ws[f"{perf}{r}"] = (f'=IF(AND(ISNUMBER({act}{r}),ISNUMBER({tgt}{r})),'
                                    f'IF({tgt}{r}>0,{miss},""),"")')
            if i >= first + band_n:
                window = f"{perf}{r - band_n}:{perf}{r - 1}"
                ws[f"{band}{r}"] = (f'=IF(COUNT({window})<{band_n},"",MAX({s["band_floor"]},MIN({s["band_cap"]},'
                                    f'QUARTILE({window},3)-QUARTILE({window},1))))')
            for col, fmt in [(act, d.NUMBER_FORMAT[kpi.fmt]), (tgt, d.NUMBER_FORMAT[kpi.fmt]),
                             (perf, d.PERF_FORMAT), (band, "0.0%")]:
                ws[f"{col}{r}"].number_format = fmt
                ws[f"{col}{r}"].font = _font()
    ws.column_dimensions["A"].width = 10
    for c in range(2, 2 + 4 * len(KPIS)):
        ws.column_dimensions[get_column_letter(c)].width = 11
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B3"
    return out


# --- report ---------------------------------------------------------------------

def _report_sheet(ws, grain: str, data: PeriodData, kpi_cols: dict[str, dict[str, str]], last_action_row: int,
                  default_key: str, cfg: Config) -> None:
    _, kpi_tab, data_tab = _sheet_names(grain)
    kpi_q, data_q = _q(kpi_tab), _q(data_tab)
    last = FIRST_ROW + len(data.table) - 1
    lag = WINDOWS[grain]["lag"]
    word = PERIOD_WORD[grain]

    def col(letter: str) -> str:
        return f"{kpi_q}!${letter}${FIRST_ROW}:${letter}${last}"

    def pick(letter: str, offset: str = "") -> str:
        return f"INDEX({col(letter)},{ROW_CELL}{offset})"

    ws.sheet_view.showGridLines = False
    ws["B2"] = f"{GRAIN_LABEL[grain]} KPI report"
    ws["B2"].font = _font(bold=True, size=16)
    ws["B3"] = (f"{cfg.brand} (synthetic data). Every target is a formula on the brand's own history; "
                f"click any figure to trace it.")
    ws["B3"].font = _font(color=d.INK_2)

    ws["B5"] = word
    ws["B5"].font = _font(bold=True)
    ws["C5"] = default_key
    ws["C5"].font = _font(bold=True, color=d.INPUT_BLUE)
    ws["C5"].fill = _fill(d.SETTING_FILL)
    ws["C5"].border = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)
    ws["C5"].alignment = Alignment(horizontal="center")
    first_targeted = FIRST_ROW + first_target_index(data.table, next(k for k in KPIS if k.id == "net_rev"), grain)
    # no leading "=": the file format stores list sources bare, and Google Sheets drops the rule otherwise
    dv = DataValidation(type="list", formula1=f"{kpi_q}!$A${first_targeted}:$A${last}", allow_blank=False)
    dv.error = f"Pick a {word.lower()} from the list."
    dv.prompt = f"Pick any {word.lower()} with a target"
    ws.add_data_validation(dv)
    dv.add("C5")
    ws[ROW_CELL.replace("$", "")] = f"=MATCH($C$5,{col('A')},0)"
    for cell, column in [("D5", "B"), ("F5", "C")]:
        ws[cell] = f"=INDEX({data_q}!${column}${FIRST_ROW}:${column}${last},{ROW_CELL})"
        ws[cell].number_format = "d mmm yyyy"
        ws[cell].font = _font()
    ws["E5"] = "to"
    ws["E5"].font = _font(color=d.MUTED)
    ws["E5"].alignment = Alignment(horizontal="center")
    ws["G5"] = f"Pick any {word.lower()} in the yellow cell."
    ws["G5"].font = _font(color=d.MUTED, italic=True)
    ws["B6"] = ("On track: the miss is within the KPI's normal swing.  Watch: up to twice its normal swing.  "
                "Off track: more than twice.  Each KPI's normal swing comes from its own past misses (see Read me).")
    ws["B6"].font = _font(size=9, color=d.INK_2)

    _trend_charts(ws, grain, data, kpi_cols, col)

    r = SECTION_START
    status_cells = []
    for section in SECTIONS:
        headers = [section, "Actual", "Target", "vs target", "Status", "Normal swing", "Last year", "Note"]
        for j, h in enumerate(headers, start=2):
            ws.cell(r, j, h)
        _style_header(ws[r][1:9])
        ws.cell(r, 2).font = _font(bold=True, size=11)
        r += 1
        for kpi in [k for k in KPIS if k.section == section]:
            c = kpi_cols[kpi.id]
            ws[f"B{r}"] = kpi.label
            ws[f"C{r}"] = f"={pick(c['actual'])}"
            ws[f"D{r}"] = f"={pick(c['target'])}"
            ws[f"E{r}"] = f"={pick(c['perf'])}"
            ws[f"G{r}"] = f'=IF(ISNUMBER(E{r}),Y{r}/2,"")'  # half the band: the miss this KPI usually shows
            ws[f"H{r}"] = f'=IF({ROW_CELL}>{lag},{pick(c["actual"], f"-{lag}")},"")'
            ws[f"I{r}"] = NOTES.get(kpi.id, "")
            ws[f"Y{r}"] = (f'=IF(ISNUMBER({pick(c["band"])}),{pick(c["band"])},'
                           f'Settings!$B${SETTING_ROW["default_band"]})')
            ws[f"Z{r}"] = kpi.direction
            ws[f"F{r}"] = (f'=IF(ISNUMBER(E{r}),IF($Z{r}="neutral",IF(E{r}>Y{r},"Above plan",'
                           f'IF(E{r}<-Y{r},"Below plan","On plan")),IF(E{r}>=-Y{r}/2,"On track",'
                           f'IF(E{r}>=-Y{r},"Watch","Off track"))),"n/a")')
            for letter, fmt in [("C", d.NUMBER_FORMAT[kpi.fmt]), ("D", d.NUMBER_FORMAT[kpi.fmt]),
                                ("E", d.PERF_FORMAT), ("G", SWING_FORMAT), ("H", d.NUMBER_FORMAT[kpi.fmt])]:
                ws[f"{letter}{r}"].number_format = fmt
            for letter in "BCDEFGHI":
                ws[f"{letter}{r}"].font = _font(color=d.INK_2 if letter in "GI" else d.INK)
                ws[f"{letter}{r}"].border = BOTTOM
            ws[f"F{r}"].alignment = Alignment(horizontal="center")
            status_cells.append(f"F{r}")
            r += 1
        ws[f"B{r}"] = "What to do"
        ws[f"B{r}"].font = _font(bold=True, color=d.INK_2)
        r += 1
        for n in range(1, 4):
            key = f'"{grain}|"&$C$5&"|{section}|{n}"'
            ws[f"B{r}"] = (f'=IFERROR("• "&INDEX(Actions!$F$2:$F${last_action_row},'
                           f'MATCH({key},Actions!$A$2:$A${last_action_row},0)),"")')
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
            ws[f"B{r}"].alignment = WRAP_TOP
            ws[f"B{r}"].font = _font()
            ws.row_dimensions[r].height = 28
            r += 1
        r += 1

    rng = f"F{SECTION_START}:F{r}"
    for label, (_, tint) in d.STATUS.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{label}"'],
                                                      fill=_fill(tint), font=Font(name=d.FONT, color=d.INK)))
    for letter, width in zip("ABCDEFGHI", [2, 30, 14, 14, 11, 12, 13, 14, 40]):
        ws.column_dimensions[letter].width = width
    for letter in ["X", "Y", "Z"]:
        ws.column_dimensions[letter].hidden = True
    ws.freeze_panes = "A7"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _axis_floor(values: pd.Series, money: bool) -> float:
    """A round number a bit under the lowest value, so the axis doesn't flatten small changes against zero."""
    low = float(values.min()) * 0.85
    if low <= 0:
        return 0.0
    step = 10 ** math.floor(math.log10(low)) / 2 if money else 0.0025
    return math.floor(low / step) * step


def _trend_charts(ws, grain: str, data: PeriodData, kpi_cols: dict[str, dict[str, str]], col) -> None:
    """Two line charts on the report, fed by this report's block on the Chart data tab.

    The chart data lives on its own visible tab: spreadsheet charts skip hidden
    cells by default, and Google Sheets drops such series on import.
    """
    targeted = data.table.iloc[first_target_index(data.table, next(k for k in KPIS if k.id == "net_rev"), grain):]
    floors = {"net": _axis_floor(targeted["net_rev"], money=True),
              "cr": _axis_floor(targeted["orders"] / targeted["sessions"], money=False)}
    data_ws = ws.parent["Chart data"]
    row_cell = f"{_q(ws.title)}!{ROW_CELL}"
    n = CHART_PERIODS[grain]
    c0 = CHART_FIRST_COL[grain]
    head = 3
    unit = {"week": "weeks", "month": "months", "quarter": "quarters"}[grain]
    data_ws.cell(1, c0, f"{ws.title}: last {n} {unit} up to the chosen one").font = _font(bold=True)
    data_ws.cell(2, c0, "#N/A marks a period with no figure yet, so the chart skips it.").font = \
        _font(color=d.MUTED, italic=True)
    series = [(None, "Period", None), (("net_rev", "actual"), "Net revenue", d.NUMBER_FORMAT["money"]),
              (("net_rev", "target"), "Target", d.NUMBER_FORMAT["money"]),
              (("cr", "actual"), "Conversion rate", d.NUMBER_FORMAT["pct2"]),
              (("cr", "target"), "Target", d.NUMBER_FORMAT["pct2"])]
    for j, (source, title, fmt) in enumerate(series):
        letter = get_column_letter(c0 + j)
        data_ws[f"{letter}{head}"] = title
        _style_header([data_ws[f"{letter}{head}"]])
        data_ws.column_dimensions[letter].width = 13
        for k in range(n):
            back = n - 1 - k
            column = "A" if source is None else kpi_cols[source[0]][source[1]]
            value = f"INDEX({col(column)},{row_cell}{f'-{back}' if back else ''})"
            missing = '""' if source is None else "NA()"  # Excel plots "" as zero; it skips #N/A
            shown = value if source is None else f"IF(ISNUMBER({value}),{value},{missing})"
            cell = data_ws[f"{letter}{head + 1 + k}"]
            cell.value = f"=IF({row_cell}>{back},{shown},{missing})"
            cell.font = _font()
            if fmt:
                cell.number_format = fmt
    cats = Reference(data_ws, min_col=c0, min_row=head + 1, max_row=head + n)
    charts = [("Net revenue vs target", c0 + 1, '"€"#,##0', "B8", 12.6, floors["net"]),
              ("Conversion rate vs target", c0 + 3, "0.0%", "F8", 14.6, floors["cr"])]
    for title, first_col, fmt, anchor, width, floor in charts:
        chart = LineChart()
        chart.title = f"{title}, last {n} {unit}"
        chart.title.overlay = False  # openpyxl leaves this unset, and Excel then draws the title over the plot
        chart.height, chart.width = 7.4, width
        chart.y_axis.scaling.min = floor
        chart.add_data(Reference(data_ws, min_col=first_col, max_col=first_col + 1, min_row=head, max_row=head + n),
                       titles_from_data=True)
        chart.set_categories(cats)
        chart.visible_cells_only = False
        for s, color, line_width in zip(chart.series, [d.ACCENT, d.MUTED], [28575, 19050]):
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = line_width
            s.smooth = False
            s.marker = Marker(symbol="none")
        chart.y_axis.number_format = fmt
        chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=d.GRID))
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.legend.position = "b"
        chart.legend.overlay = False  # same for the legend, which otherwise covers the period labels
        ws.add_chart(chart, anchor)


# --- reference tabs -------------------------------------------------------------

def _definitions(ws) -> None:
    headers = ["Section", "KPI", "Definition", "Calculation", "Better when"]
    for j, h in enumerate(headers, start=1):
        ws.cell(1, j, h)
    _style_header(ws[1][:5])
    label = {c: h for c, h, _, _ in BASE_COLUMNS}
    better = {"up": "Higher", "down": "Lower", "neutral": "Neutral (pacing)"}
    for i, k in enumerate(KPIS, start=2):
        calc = label[k.num] if not k.den else f"{label[k.num]} ÷ {label[k.den]} (ratio of sums)"
        for j, v in enumerate([k.section, k.label, k.definition, calc, better[k.direction]], start=1):
            cell = ws.cell(i, j, v)
            cell.font = _font(bold=j == 2)
            cell.alignment = WRAP_TOP
    for letter, width in zip("ABCDE", [18, 26, 70, 52, 16]):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"


def _checks(ws, checks: DataChecks) -> None:
    for j, h in enumerate(["Source", "Check", "Result"], start=1):
        ws.cell(1, j, h)
    _style_header(ws[1][:3])
    for i, (source, check, result) in enumerate(checks.rows, start=2):
        for j, v in enumerate([source, check, result], start=1):
            cell = ws.cell(i, j, v)
            cell.font = _font(bold=j == 2)
            cell.alignment = WRAP_TOP
    for letter, width in zip("ABC", [18, 32, 110]):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"


def _read_me(ws, cfg: Config, defaults: dict[str, str]) -> None:
    ws.sheet_view.showGridLines = False
    lines = [
        ("title", f"{cfg.brand} KPI workbook"),
        ("sub", f"Built with Claude Code from raw Shopify, Meta Ads, Google Ads, TikTok Ads and GA4 exports. "
                f"The data is synthetic. Generated {datetime.now():%d %b %Y}; data to {cfg.as_of:%d %b %Y}."),
        ("head", "How to use it"),
        ("text", f"Open a report tab and pick a period in the yellow cell at the top. Defaults: "
                 f"{defaults['week']}, {defaults['month']} and {defaults['quarter']}."),
        ("text", "Each section lists its KPIs against target, then what to do about them."),
        ("head", "How targets are set"),
        ("text", "Every target comes from the brand's own numbers; none is typed in. A target is the same period "
                 "last year (for weeks, the average of that week and its neighbours), scaled by the brand's current "
                 "growth: the last 13 weeks (3 months, 2 quarters) against the same periods a year earlier. Growth is "
                 "capped at +50% and -33% so one odd period can't set a silly target."),
        ("text", "Ratio KPIs such as conversion rate, MER and CAC use ratios of sums, so a big week counts more "
                 "than a small one."),
        ("head", "How status is decided"),
        ("text", "Each KPI misses its target by a different amount in a normal period: revenue might swing ±5% "
                 "week to week while conversion rate barely moves. A KPI's normal swing is half the spread "
                 "(interquartile range) of its own misses over the previous periods, shown in the report's Normal "
                 "swing column. A miss within the normal swing is On track, up to twice the swing is Watch, and "
                 "beyond that is Off track. So a 4% revenue miss can be normal while a 4% conversion miss is not."),
        ("head", "Colour code"),
        ("text", "Blue numbers come from the exports. Black numbers are formulas. Yellow cells are settings you "
                 "can change (period pickers, Settings tab)."),
        ("head", "Tabs"),
        ("text", "Weekly, Monthly and Quarterly report: the reports. KPIs tabs: actual, target, miss and band for "
                 "every period, all formulas. Data tabs: the brand's numbers per period. Chart data: the figures "
                 "behind the report charts, following each report's chosen period. Actions: the "
                 "recommendations behind each report section, written by rules from the numbers. KPI definitions, "
                 "Data checks (everything the cleaning step fixed or flagged), Settings."),
        ("head", "Cost assumptions"),
        ("text", f"Prices include {cfg.vat_rate:.0%} VAT; all revenue is shown ex VAT. Payment fees "
                 f"{cfg.payment_fee_pct:.1%} + €{cfg.payment_fee_fixed:.2f} per order. Fulfilment "
                 f"€{cfg.fulfilment_per_order:.2f} per order + €{cfg.fulfilment_per_extra_unit:.2f} per extra unit. "
                 f"A SKU missing from the cost sheet is costed at {cfg.missing_cost_share:.0%} of its net price."),
    ]
    r = 2
    for kind, text in lines:
        if kind == "head":
            r += 1  # a blank row above each heading
        cell = ws.cell(r, 2, text)
        if kind == "title":
            cell.font = _font(bold=True, size=16)
        elif kind == "sub":
            cell.font = _font(color=d.INK_2)
        elif kind == "head":
            cell.font = _font(bold=True, size=11)
        else:
            cell.font = _font()
            cell.alignment = WRAP_TOP
            ws.row_dimensions[r].height = 15 * max(1, len(text) // 105 + 1)
        r += 1
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 110
